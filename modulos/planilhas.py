"""Leitura de QUALQUER aba das planilhas da pasta dados.

O MAX usa este módulo para responder perguntas sobre qualquer aba
(solicitações de compra, apontamentos, base de vendas, fornecedores,
estoque e carteira...), não só a análise de DOH.

Como cada aba tem um formato, a leitura é automática:
- o cabeçalho é a linha, entre as primeiras, com mais textos/datas;
- nas abas de ANÁLISE (Fiesta, Talher...), a tabela de produtos começa
  na linha "Cod Prodt" e o quadro do topo vira uma tabela à parte,
  "<aba> - RESUMO POR SEGMENTO";
- linhas e colunas totalmente vazias são descartadas.

Quem calcula é sempre o Python (filtros, somas, médias, contagens), com
os valores da última versão SALVA de cada planilha. Ela só é lida de
novo quando o arquivo muda.
"""

import math
import re
import threading
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


PASTA_DADOS = Path(__file__).resolve().parent.parent / "dados"

LINHAS_VAZIAS_PARA_PARAR = 500   # algumas abas vão "formatadas" até a linha 1.048.576
LINHAS_PROCURA_CABECALHO = 15
LIMITE_LINHAS = 100
LIMITE_COLUNAS = 30
LIMITE_GRUPOS = 60
LIMITE_TEXTO = 120


# ============================================================
# TEXTO
# ============================================================

def normalizar(texto):
    """Maiúsculas, sem acentos e com espaços simples."""
    texto = unicodedata.normalize("NFD", str(texto))
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto).strip().upper()


def _radical(palavra):
    """Tira o plural simples: "COPOS" -> "COPO", "TALHERES" -> "TALHER"."""
    if len(palavra) > 4 and palavra.endswith("ES"):
        return palavra[:-2]
    if len(palavra) > 3 and palavra.endswith("S"):
        return palavra[:-1]
    return palavra


def _vazio(valor):
    if valor is None:
        return True
    if isinstance(valor, float) and math.isnan(valor):
        return True
    return isinstance(valor, str) and not valor.strip()


# ============================================================
# LEITURA DAS ABAS
# ============================================================

@dataclass
class Tabela:
    nome: str
    arquivo: str
    oculta: bool
    df: pd.DataFrame = field(repr=False)


def _nome_coluna(valor, indice):
    """Texto do cabeçalho; datas viram "08/2025" (mês) ou "28/10/2026"."""
    if isinstance(valor, datetime):
        if valor.day == 1 and valor.time() == time(0):
            return f"{valor:%m/%Y}"
        return f"{valor:%d/%m/%Y}"
    if isinstance(valor, date):
        return f"{valor:%d/%m/%Y}"
    if _vazio(valor):
        letra = ""
        n = indice + 1
        while n:
            n, resto = divmod(n - 1, 26)
            letra = chr(65 + resto) + letra
        return f"Coluna {letra}"
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", " ", str(valor)).strip()


def _sem_repetir(nomes):
    vistos = {}
    saida = []
    for nome in nomes:
        chave = normalizar(nome)
        if chave in vistos:
            vistos[chave] += 1
            nome = f"{nome} ({vistos[chave]})"
        else:
            vistos[chave] = 1
        saida.append(nome)
    return saida


def _ler_linhas(ws):
    """Todas as linhas com algum valor (para depois de muitas vazias seguidas)."""
    linhas = []
    vazias = 0
    for linha in ws.iter_rows(values_only=True):
        if all(_vazio(v) for v in linha):
            vazias += 1
            if vazias >= LINHAS_VAZIAS_PARA_PARAR:
                break
            linhas.append(None)
        else:
            vazias = 0
            linhas.append(list(linha))
    while linhas and linhas[-1] is None:
        linhas.pop()
    largura = max((len(l) for l in linhas if l), default=0)
    return [
        (l + [None] * (largura - len(l))) if l else [None] * largura
        for l in linhas
    ]


def _montar_df(cabecalho, linhas, manter=None):
    """DataFrame só com as colunas que têm nome ou algum dado."""
    nomes = _sem_repetir([_nome_coluna(v, i) for i, v in enumerate(cabecalho)])
    df = pd.DataFrame(linhas, columns=nomes, dtype=object)
    preenchidas = pd.Series(
        [not all(_vazio(v) for v in linha) for linha in df.itertuples(index=False)],
        index=df.index, dtype=bool,
    )
    df = df.loc[preenchidas]
    colunas = [
        nome for i, nome in enumerate(nomes)
        if (manter and nome in manter)
        or not _vazio(cabecalho[i])
        or not df[nome].map(_vazio).all()
    ]
    return df[colunas].reset_index(drop=True)


def _linha_do_cabecalho(linhas):
    """Entre as primeiras linhas, a que tem mais textos e datas."""
    melhor, pontos_melhor = 0, -1
    for i, linha in enumerate(linhas[:LINHAS_PROCURA_CABECALHO]):
        if not linha:
            continue
        pontos = sum(
            1 for v in linha
            if isinstance(v, (str, datetime, date)) and not _vazio(v)
        )
        if pontos > pontos_melhor:
            melhor, pontos_melhor = i, pontos
    return melhor


def _tabelas_da_aba(ws, arquivo):
    linhas = _ler_linhas(ws)
    oculta = ws.sheet_state != "visible"
    if not linhas:
        return []

    # Abas de análise: tabela de produtos a partir de "Cod Prodt".
    produtos = None
    for i, linha in enumerate(linhas[:80]):
        if linha and any(normalizar(v) in ("COD PRODT", "CODIGO") for v in linha if isinstance(v, str)):
            if i > 2:
                produtos = i
            break

    if produtos is None:
        h = _linha_do_cabecalho(linhas)
        dados = [l for l in linhas[h + 1:] if l]
        return [Tabela(ws.title, arquivo, oculta, _montar_df(linhas[h], dados))]

    tabelas = []
    cab = linhas[produtos]
    col_codigo = next(
        j for j, v in enumerate(cab)
        if isinstance(v, str) and normalizar(v) in ("COD PRODT", "CODIGO")
    )
    dados = [l for l in linhas[produtos + 1:] if l and not _vazio(l[col_codigo])]
    tabelas.append(Tabela(ws.title, arquivo, oculta, _montar_df(cab, dados)))

    # Quadro do topo: linha 1 = nomes, linha 2 = unidades, depois segmentos.
    topo, unidades = linhas[0], linhas[1] if len(linhas) > 1 else None
    if topo:
        col_seg = next((j for j, v in enumerate(topo) if not _vazio(v)), 1) - 1
        if col_seg >= 0:
            nomes = []
            for j, v in enumerate(topo):
                if j < col_seg:
                    nomes.append(None)
                elif j == col_seg:
                    nomes.append("Segmento")
                else:
                    unidade = unidades[j] if unidades and isinstance(unidades[j], str) else ""
                    nome = _nome_coluna(v, j)
                    if unidade.strip() and not normalizar(unidade).startswith("CLIQUE"):
                        nome = f"{nome} ({unidade.strip()})"
                    nomes.append(nome if not _vazio(v) else None)
            segs = [
                l for l in linhas[2:produtos]
                if l and not _vazio(l[col_seg])
            ]
            if segs:
                df = _montar_df(nomes, segs, manter={"Segmento"})
                df = df[[c for c in df.columns if not c.startswith("Coluna ")]]
                tabelas.append(Tabela(
                    f"{ws.title} - RESUMO POR SEGMENTO", arquivo, oculta, df
                ))
    return tabelas


_cache = {}
_trava = threading.Lock()


def arquivos():
    """Planilhas da pasta dados (sem os temporários "~$" do Excel)."""
    return sorted(
        p for p in PASTA_DADOS.glob("*.xls[xm]")
        if not p.name.startswith("~$")
    )


def _tabelas_do_arquivo(caminho):
    versao = caminho.stat().st_mtime
    with _trava:
        item = _cache.get(caminho)
        if item and item[0] == versao:
            return item[1]
        wb = load_workbook(caminho, read_only=True, data_only=True)
        try:
            tabelas = []
            for ws in wb.worksheets:
                try:
                    tabelas.extend(_tabelas_da_aba(ws, caminho.name))
                except Exception as erro:  # uma aba estranha não derruba as outras
                    print(f"[MAX] Não consegui ler a aba {ws.title!r}: {erro}", flush=True)
        finally:
            wb.close()
        _cache[caminho] = (versao, tabelas)
        return tabelas


def tabelas():
    """Todas as tabelas de todas as planilhas."""
    todas = []
    for caminho in arquivos():
        todas.extend(_tabelas_do_arquivo(caminho))
    # Mesmo nome de aba em arquivos diferentes: acrescenta o arquivo.
    contagem = {}
    for t in todas:
        contagem[normalizar(t.nome)] = contagem.get(normalizar(t.nome), 0) + 1
    for t in todas:
        if contagem[normalizar(t.nome)] > 1 and "[" not in t.nome:
            t.nome = f"{t.nome} [{t.arquivo}]"
    return todas


def resolver_tabela(nome):
    """Acha a tabela pelo nome completo ou por parte dele.

    Devolve (Tabela, None) ou (None, dicionário de erro).
    """
    todas = tabelas()
    alvo = normalizar(nome)
    if not alvo:
        return None, {"erro": "Informe qual aba usar.", "abas": [t.nome for t in todas]}

    exatas = [t for t in todas if normalizar(t.nome) == alvo]
    if len(exatas) == 1:
        return exatas[0], None

    palavras = [_radical(p) for p in re.split(r"[\s\-_.]+", alvo) if p and p != "ANALISE"]
    candidatas = [t for t in todas if all(p in normalizar(t.nome) for p in palavras)]
    if not any(p in ("RESUMO", "SEGMENTO") for p in palavras):
        sem_resumo = [t for t in candidatas if "RESUMO POR SEGMENTO" not in normalizar(t.nome)]
        candidatas = sem_resumo or candidatas

    if len(candidatas) == 1:
        return candidatas[0], None
    if not candidatas:
        return None, {
            "erro": f"Não encontrei a aba '{nome}'.",
            "abas": [t.nome for t in todas],
        }
    return None, {
        "erro": f"'{nome}' combina com mais de uma aba. Pergunte ao usuário qual delas.",
        "opcoes": [t.nome for t in candidatas],
    }


def resolver_coluna(df, nome):
    """Acha a coluna pelo nome (sem diferenciar acentos e maiúsculas)."""
    alvo = normalizar(nome).strip(" .:")
    colunas = list(df.columns)
    for criterio in (
        lambda c: normalizar(c) == alvo,
        lambda c: normalizar(c).strip(" .:") == alvo,
        lambda c: normalizar(c).startswith(alvo),
        lambda c: alvo in normalizar(c),
    ):
        achadas = [c for c in colunas if criterio(c)]
        if len(achadas) == 1:
            return achadas[0], None
        if len(achadas) > 1:
            return None, {
                "erro": f"A coluna '{nome}' combina com mais de uma: {achadas}.",
                "colunas": colunas,
            }
    return None, {"erro": f"Não existe a coluna '{nome}' nesta aba.", "colunas": colunas}


# ============================================================
# VALORES
# ============================================================

def valor_para_json(valor):
    """Valor da planilha em formato simples para o Gemini."""
    if _vazio(valor):
        return None
    if hasattr(valor, "item") and not isinstance(valor, (datetime, date)):
        valor = valor.item()
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, datetime):
        return f"{valor:%d/%m/%Y}" if valor.time() == time(0) else f"{valor:%d/%m/%Y %H:%M}"
    if isinstance(valor, date):
        return f"{valor:%d/%m/%Y}"
    if isinstance(valor, time):
        return f"{valor:%H:%M}"
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        if math.isinf(valor):
            return None
        if valor.is_integer():
            return int(valor)
        return round(valor, 2) if abs(valor) >= 1 else round(valor, 6)
    texto = re.sub(r"\s+", " ", str(valor)).strip()
    return texto if len(texto) <= LIMITE_TEXTO else texto[:LIMITE_TEXTO] + "…"


def _numeros(serie):
    def conv(v):
        if isinstance(v, bool) or _vazio(v):
            return math.nan
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, str):
            t = v.strip().replace("%", "")
            if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", t):   # 1.234,5
                t = t.replace(".", "").replace(",", ".")
            elif re.fullmatch(r"-?\d+,\d+", t):                 # 12,5
                t = t.replace(",", ".")
            try:
                return float(t)
            except ValueError:
                return math.nan
        return math.nan
    return serie.map(conv).astype(float)


def _datas(serie):
    def conv(v):
        if isinstance(v, datetime):
            return pd.Timestamp(v)
        if isinstance(v, date):
            return pd.Timestamp(v)
        if isinstance(v, str):
            return _data_texto(v)
        return pd.NaT
    return pd.to_datetime(serie.map(conv), errors="coerce")


def _data_texto(texto):
    t = str(texto).strip()
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%d/%m/%Y %H:%M", "%Y-%m-%d", "%m/%Y"):
        try:
            return pd.Timestamp(datetime.strptime(t, formato))
        except ValueError:
            pass
    return pd.NaT


def _texto(serie):
    return serie.map(lambda v: "" if _vazio(v) else normalizar(valor_para_json(v)))


def _tipo(serie):
    """'numero', 'data' ou 'texto', pelo que predomina na coluna."""
    preenchidos = serie[~serie.map(_vazio)]
    if preenchidos.empty:
        return "texto"
    if _datas(preenchidos).notna().mean() >= 0.8 and preenchidos.map(
        lambda v: isinstance(v, (datetime, date))
    ).any():
        return "data"
    if _numeros(preenchidos).notna().mean() >= 0.8:
        return "numero"
    return "texto"


# ============================================================
# FILTROS
# ============================================================

_OPERADORES_PALAVRA = [
    (r"n[aã]o\s+cont[eé]m", "nao_contem"),
    (r"cont[eé]m", "contem"),
    (r"come[cç]a\s+com", "comeca"),
    (r"termina\s+com", "termina"),
    (r"n[aã]o\s+(?:est[aá]\s+|[eé]\s+)?vazi[oa]", "nao_vazio"),
    (r"(?:est[aá]\s+|[eé]\s+)?vazi[oa]", "vazio"),
    (r"entre", "entre"),
]
_OPERADORES_SIMBOLO = [(">=", ">="), ("<=", "<="), ("!=", "!="), ("<>", "!="), ("=", "="), (">", ">"), ("<", "<")]


def _separar_condicao(condicao):
    """'Status = AGUARDANDO' -> ('Status', '=', 'AGUARDANDO')."""
    for padrao, op in _OPERADORES_PALAVRA:
        m = re.match(rf"^(.+?)\s+{padrao}(?:\s+(.*))?$", condicao, re.IGNORECASE)
        if m:
            return m.group(1).strip(), op, (m.group(2) or "").strip()
    for simbolo, op in _OPERADORES_SIMBOLO:
        m = re.match(rf"^(.+?)\s*{re.escape(simbolo)}\s*(.*)$", condicao)
        if m:
            return m.group(1).strip(), op, m.group(2).strip()
    return None


def _sem_aspas(texto):
    return texto.strip().strip("\"'“”‘’").strip()


def aplicar_filtros(df, filtros):
    """Filtros separados por ";". Ex.: "Status = AGUARDANDO; Quantidade > 100".

    Devolve (df filtrado, None) ou (None, dicionário de erro).
    """
    for condicao in [c.strip() for c in str(filtros or "").split(";") if c.strip()]:
        partes = _separar_condicao(condicao)
        if not partes:
            return None, {
                "erro": f"Não entendi o filtro '{condicao}'.",
                "exemplo": "Status = AGUARDANDO; Quantidade > 100; Descrição contém COPO",
            }
        nome, op, valor = partes
        coluna, erro = resolver_coluna(df, nome)
        if erro:
            return None, erro
        serie = df[coluna]
        valor = _sem_aspas(valor)

        if op == "vazio":
            mascara = serie.map(_vazio)
        elif op == "nao_vazio":
            mascara = ~serie.map(_vazio)
        elif op in ("contem", "nao_contem", "comeca", "termina"):
            alvo = normalizar(valor)
            texto = _texto(serie)
            if op == "comeca":
                mascara = texto.str.startswith(alvo)
            elif op == "termina":
                mascara = texto.str.endswith(alvo)
            else:
                mascara = texto.str.contains(alvo, regex=False)
                if op == "nao_contem":
                    mascara = ~mascara
        else:
            mascara, erro = _comparar(serie, op, valor, coluna)
            if erro:
                return None, erro
        df = df[mascara.fillna(False).astype(bool)]
    return df, None


def _comparar(serie, op, valor, coluna):
    if op == "entre":
        partes = re.split(r"\s+e\s+|\s*;\s*|\s*,\s*(?=\S)", valor, maxsplit=1)
        if len(partes) != 2:
            return None, {"erro": f"Use: {coluna} entre 10 e 20."}
        a, erro = _comparar(serie, ">=", partes[0], coluna)
        if erro:
            return None, erro
        b, erro = _comparar(serie, "<=", partes[1], coluna)
        if erro:
            return None, erro
        return a & b, None

    tipo = _tipo(serie)
    alvo_data = _data_texto(valor)
    alvo_num = _numeros(pd.Series([valor])).iloc[0]

    if tipo == "data" and not pd.isna(alvo_data):
        esquerda, direita = _datas(serie), alvo_data
        if op in ("=", "!=") and alvo_data.time() == time(0):
            esquerda = esquerda.dt.normalize()
    elif not math.isnan(alvo_num) and op in (">", ">=", "<", "<=") or (
        not math.isnan(alvo_num) and tipo == "numero"
    ):
        esquerda, direita = _numeros(serie), alvo_num
    elif op in ("=", "!="):
        esquerda, direita = _texto(serie), normalizar(valor)
    else:
        return None, {
            "erro": f"Para usar '{op}' na coluna '{coluna}', informe um número ou uma data (dd/mm/aaaa)."
        }

    comparacoes = {
        "=": lambda a, b: a == b,
        "!=": lambda a, b: a != b,
        ">": lambda a, b: a > b,
        ">=": lambda a, b: a >= b,
        "<": lambda a, b: a < b,
        "<=": lambda a, b: a <= b,
    }
    return comparacoes[op](esquerda, direita), None


# ============================================================
# CONSULTAS
# ============================================================

def _ordenar(df, ordenar_por):
    texto = str(ordenar_por or "").strip()
    if not texto:
        return df, None
    decrescente = False
    m = re.match(r"^(.*?)\s+(desc|decrescente|maior para (?:o )?menor|asc|crescente|menor para (?:o )?maior)$",
                 texto, re.IGNORECASE)
    if m:
        texto = m.group(1)
        decrescente = normalizar(m.group(2)).startswith(("DESC", "DECRESC", "MAIOR"))
    coluna, erro = resolver_coluna(df, texto)
    if erro:
        return None, erro
    tipo = _tipo(df[coluna])
    chave = {"numero": _numeros, "data": _datas, "texto": _texto}[tipo](df[coluna])
    ordem = chave.sort_values(ascending=not decrescente, na_position="last").index
    return df.loc[ordem], None


def _escolher_colunas(df, colunas):
    nomes = [c.strip() for c in str(colunas or "").split(",") if c.strip()]
    if not nomes:
        return list(df.columns[:LIMITE_COLUNAS]), None
    escolhidas = []
    for nome in nomes:
        coluna, erro = resolver_coluna(df, nome)
        if erro:
            return None, erro
        if coluna not in escolhidas:
            escolhidas.append(coluna)
    return escolhidas, None


def linhas_json(df, colunas):
    return [
        {c: v for c in colunas if (v := valor_para_json(linha[c])) is not None}
        for _, linha in df.iterrows()
    ]


def catalogo():
    """Lista de abas com o tamanho e as colunas de cada uma."""
    todas = tabelas()
    varios_arquivos = len({t.arquivo for t in todas}) > 1
    saida = []
    for t in todas:
        item = {
            "aba": t.nome,
            "linhas": len(t.df),
            "colunas": [str(c) for c in t.df.columns[:40]],
        }
        if len(t.df.columns) > 40:
            item["outras_colunas"] = len(t.df.columns) - 40
        if t.oculta:
            item["oculta"] = True
        if varios_arquivos:
            item["arquivo"] = t.arquivo
        saida.append(item)
    return saida


def consultar(aba, filtros, colunas, ordenar_por, limite):
    tabela, erro = resolver_tabela(aba)
    if erro:
        return erro
    df, erro = aplicar_filtros(tabela.df, filtros)
    if erro:
        return {"aba": tabela.nome, **erro}
    df, erro = _ordenar(df, ordenar_por)
    if erro:
        return {"aba": tabela.nome, **erro}
    escolhidas, erro = _escolher_colunas(tabela.df, colunas)
    if erro:
        return {"aba": tabela.nome, **erro}

    try:
        limite = int(limite)
    except (TypeError, ValueError):
        limite = 20
    limite = max(1, min(limite, LIMITE_LINHAS))

    resultado = {
        "aba": tabela.nome,
        "linhas_encontradas": len(df),
        "linhas_mostradas": min(limite, len(df)),
        "linhas": linhas_json(df.head(limite), escolhidas),
    }
    if len(tabela.df.columns) > len(escolhidas):
        resultado["outras_colunas_da_aba"] = [c for c in tabela.df.columns if c not in escolhidas][:40]
    return resultado


_CALCULOS = {
    "SOMA": "soma", "SOMAR": "soma", "TOTAL": "soma",
    "MEDIA": "media", "MEDIO": "media",
    "MAX": "maximo", "MAXIMO": "maximo", "MAIOR": "maximo",
    "MIN": "minimo", "MINIMO": "minimo", "MENOR": "minimo",
    "CONTAGEM": "contagem", "CONTAR": "contagem", "QUANTIDADE DE LINHAS": "contagem",
    "DISTINTOS": "distintos", "CONTAGEM DISTINTA": "distintos",
}


def _ler_calculos(df, calculos):
    lista = []
    for parte in [p.strip() for p in str(calculos or "contagem").split(",") if p.strip()]:
        m = re.match(r"^(.+?)(?:\s+(?:de|do|da|dos|das)\s+(.+))?$", parte, re.IGNORECASE)
        operacao = _CALCULOS.get(normalizar(m.group(1))) if m else None
        if operacao is None:
            return None, {
                "erro": f"Não entendi o cálculo '{parte}'.",
                "exemplo": "soma de Quantidade, media de Estoque, contagem, distintos de Fornecedor",
            }
        coluna = None
        if operacao != "contagem" or m.group(2):
            if not m.group(2):
                return None, {"erro": f"Informe a coluna: '{parte} de <coluna>'."}
            coluna, erro = resolver_coluna(df, m.group(2))
            if erro:
                return None, erro
        lista.append((operacao, coluna))
    return lista, None


def _calcular(df, operacao, coluna):
    if operacao == "contagem":
        return len(df) if coluna is None else int((~df[coluna].map(_vazio)).sum())
    if operacao == "distintos":
        return int(_texto(df[coluna]).replace("", pd.NA).nunique())
    if _tipo(df[coluna]) == "data" and operacao in ("maximo", "minimo"):
        datas = _datas(df[coluna])
        valor = datas.max() if operacao == "maximo" else datas.min()
        return None if pd.isna(valor) else valor_para_json(valor.to_pydatetime())
    numeros = _numeros(df[coluna])
    valor = {
        "soma": numeros.sum,
        "media": numeros.mean,
        "maximo": numeros.max,
        "minimo": numeros.min,
    }[operacao]()
    return None if pd.isna(valor) else valor_para_json(float(valor))


def resumir(aba, filtros, agrupar_por, calculos):
    tabela, erro = resolver_tabela(aba)
    if erro:
        return erro
    df, erro = aplicar_filtros(tabela.df, filtros)
    if erro:
        return {"aba": tabela.nome, **erro}
    lista, erro = _ler_calculos(tabela.df, calculos)
    if erro:
        return {"aba": tabela.nome, **erro}

    def rotulo(operacao, coluna):
        return operacao if coluna is None else f"{operacao} de {coluna}"

    grupos_nomes = [g.strip() for g in str(agrupar_por or "").split(",") if g.strip()]
    if not grupos_nomes:
        return {
            "aba": tabela.nome,
            "linhas_consideradas": len(df),
            "resultado": {rotulo(o, c): _calcular(df, o, c) for o, c in lista},
        }

    grupos = []
    for nome in grupos_nomes:
        coluna, erro = resolver_coluna(tabela.df, nome)
        if erro:
            return {"aba": tabela.nome, **erro}
        grupos.append(coluna)

    chaves = df[grupos].apply(
        lambda s: s.map(lambda v: "(vazio)" if _vazio(v) else valor_para_json(v))
    )
    linhas = []
    for valores, indices in chaves.groupby(grupos, sort=False).groups.items():
        valores = valores if isinstance(valores, tuple) else (valores,)
        parte = df.loc[indices]
        item = dict(zip(grupos, valores))
        for o, c in lista:
            item[rotulo(o, c)] = _calcular(parte, o, c)
        linhas.append(item)

    primeira = rotulo(*lista[0])
    linhas.sort(key=lambda i: (i[primeira] is None, -(i[primeira] or 0) if isinstance(i[primeira], (int, float)) else 0))
    return {
        "aba": tabela.nome,
        "linhas_consideradas": len(df),
        "grupos_encontrados": len(linhas),
        "grupos": linhas[:LIMITE_GRUPOS],
    }


def buscar(termo, limite_por_aba=5):
    """Procura um texto (código, descrição, fornecedor...) em todas as abas."""
    alvo = normalizar(termo)
    if not alvo:
        return {"erro": "Informe o que procurar."}
    achados = []
    for t in tabelas():
        if t.df.empty:
            continue
        texto = t.df.apply(_texto)
        mascara = texto.apply(lambda s: s.str.contains(alvo, regex=False)).any(axis=1)
        total = int(mascara.sum())
        if not total:
            continue
        colunas = list(t.df.columns[:20])
        achados.append({
            "aba": t.nome,
            "linhas_encontradas": total,
            "exemplos": linhas_json(t.df[mascara].head(limite_por_aba), colunas),
        })
    if not achados:
        return {"termo": termo, "resultado": "Não encontrei em nenhuma aba."}
    achados.sort(key=lambda a: -a["linhas_encontradas"])
    return {"termo": termo, "abas": achados[:12]}
